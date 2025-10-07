
import sys, os, threading, math
from collections import defaultdict, OrderedDict
from PyQt5 import QtCore, QtGui, QtWidgets
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QFileDialog, QMessageBox, QTableWidgetItem
)
from docx import Document
from openpyxl import Workbook
import csv

MOCS = [5, 10, 15, 20]

class Worker(QtCore.QThread):
    progress = QtCore.pyqtSignal(str)
    finished = QtCore.pyqtSignal(dict, str)  # dict of results, csv_text

    def __init__(self, folder):
        super().__init__()
        self.folder = folder

    def run(self):
        try:
            self.progress.emit("Scanning folder...")
            files = [f for f in os.listdir(self.folder) if f.lower().endswith(".docx")]
            files.sort(key=lambda s: s.lower())
            if not files:
                self.finished.emit({}, "NO_DOCX_FOUND")
                return
            sel_files = files[-30:] if len(files) > 30 else files
            file_entries = []  # list of list of 3-digit strings per file
            idx = 0
            for fname in sel_files:
                idx += 1
                self.progress.emit(f"Reading {fname} ({idx}/{len(sel_files)})...")
                fpath = os.path.join(self.folder, fname)
                items = []
                try:
                    doc = Document(fpath)
                    for p in doc.paragraphs:
                        txt = p.text.strip()
                        digits = keep_digits(txt)
                        if len(digits) >= 3:
                            items.append(digits[-3:])
                except Exception as ex:
                    # skip file read error but report
                    self.progress.emit(f"Warning: cannot read {fname}: {ex}")
                file_entries.append(items)

            # prepare dicts
            dicts = [defaultdict(int) for _ in MOCS]
            total_entries = [0 for _ in MOCS]
            for i, moc in enumerate(MOCS):
                files_to_consider = min(moc, len(file_entries))
                start = len(file_entries) - files_to_consider
                for fi in range(start, len(file_entries)):
                    for s3 in file_entries[fi]:
                        if len(s3) == 3:
                            key = khoa_chuan3(s3)
                            dicts[i][key] += 1
                            total_entries[i] += 1

            # build output CSV-like text and results structure
            lines = []
            lines.append("PHAN TICH DAO 6 VONG - TP.HCM (GOM NHOM 6 HOAN VI)")
            top_lists = []
            for i, moc in enumerate(MOCS):
                lines.append("")
                lines.append(f"MOC {moc} KY")
                lines.append("NhomDao,SoLan,TyLe(%)")
                entries = sorted(dicts[i].items(), key=lambda kv: kv[1], reverse=True)
                if not entries:
                    lines.append("(khong co du lieu)")
                    top_lists.append([])
                else:
                    topN = min(3, len(entries))
                    lst = []
                    for r in range(topN):
                        key, cnt = entries[r]
                        rate = (cnt / total_entries[i]) if total_entries[i] > 0 else 0.0
                        lines.append(f"{key},{cnt},{rate:.2%}")
                        lst.append(key)
                    top_lists.append(lst)

            # union unique keys
            union_keys = []
            for lst in top_lists:
                for k in lst:
                    if k not in union_keys:
                        union_keys.append(k)
            if not union_keys:
                lines.append("")
                lines.append("Khong co du lieu 3-chu-so de phan tich.")
                csv_text = "\n".join(lines)
                self.finished.emit({}, csv_text)
                return

            # summary table header
            lines.append("")
            lines.append("BANG TONG HOP SO SANH CAC NHOM (TOP3 cac moc)")
            header = ["NhomDao"]
            for moc in MOCS:
                header += [f"{moc}ky_SoLan", f"{moc}ky_TyLe"]
            lines.append(",".join(header))
            results = []  # list of rows for QTableWidget
            for key in union_keys:
                row = [key]
                for i in range(len(MOCS)):
                    val = dicts[i].get(key, 0)
                    rate = (val / total_entries[i]) if total_entries[i] > 0 else 0.0
                    row += [str(val), f"{rate:.2%}"]
                lines.append(",".join(row))
                results.append(row)

            csv_text = "\n".join(lines)
            # emit result dict: headers, rows
            out = {
                "top_lists": top_lists,
                "summary_rows": results,
                "mocs": MOCS
            }
            self.finished.emit(out, csv_text)
        except Exception as ex:
            self.progress.emit(f"Error: {ex}")
            self.finished.emit({}, "ERROR")

def keep_digits(s: str) -> str:
    return "".join(ch for ch in s if ch.isdigit())

def khoa_chuan3(s: str) -> str:
    if len(s) < 3:
        return s
    a = list(s[-3:])
    a.sort()
    return "".join(a)

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("PhanTichDAO - Windows (PyQt5)")
        self.resize(900, 600)
        self.folder = ""
        self.csv_text = ""
        self.setup_ui()

    def setup_ui(self):
        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        layout = QtWidgets.QVBoxLayout(central)

        hl = QtWidgets.QHBoxLayout()
        self.btn_pick = QtWidgets.QPushButton("Chọn thư mục chứa .docx")
        self.btn_pick.clicked.connect(self.on_pick)
        hl.addWidget(self.btn_pick)
        self.btn_run = QtWidgets.QPushButton("Phân tích dữ liệu")
        self.btn_run.clicked.connect(self.on_run)
        self.btn_run.setEnabled(False)
        hl.addWidget(self.btn_run)
        self.btn_csv = QtWidgets.QPushButton("Xuất CSV")
        self.btn_csv.clicked.connect(self.on_export_csv)
        self.btn_csv.setEnabled(False)
        hl.addWidget(self.btn_csv)
        self.btn_xlsx = QtWidgets.QPushButton("Xuất Excel (.xlsx)")
        self.btn_xlsx.clicked.connect(self.on_export_xlsx)
        self.btn_xlsx.setEnabled(False)
        hl.addWidget(self.btn_xlsx)
        layout.addLayout(hl)

        self.progress = QtWidgets.QProgressBar()
        self.progress.setRange(0,0)
        self.progress.setVisible(False)
        layout.addWidget(self.progress)

        self.table = QtWidgets.QTableWidget()
        self.table.setColumnCount(1 + len(MOCS)*2)
        headers = ["NhomDao"]
        for moc in MOCS:
            headers += [f"{moc}ky_SoLan", f"{moc}ky_TyLe"]
        self.table.setHorizontalHeaderLabels(headers)
        layout.addWidget(self.table)

        self.log = QtWidgets.QTextEdit()
        self.log.setReadOnly(True)
        layout.addWidget(self.log)

    def on_pick(self):
        folder = QFileDialog.getExistingDirectory(self, "Chọn thư mục chứa các file .docx")
        if folder:
            self.folder = folder
            self.log.append(f"Đã chọn: {folder}")
            self.btn_run.setEnabled(True)

    def on_run(self):
        if not self.folder:
            QMessageBox.warning(self, "Chưa chọn thư mục", "Vui lòng chọn thư mục chứa file .docx trước.")
            return
        self.progress.setVisible(True)
        self.log.append("Bắt đầu phân tích...")
        self.worker = Worker(self.folder)
        self.worker.progress.connect(self.on_progress)
        self.worker.finished.connect(self.on_finished)
        self.worker.start()

    def on_progress(self, msg):
        self.log.append(msg)

    def on_finished(self, result, csv_text):
        self.progress.setVisible(False)
        if csv_text == "NO_DOCX_FOUND":
            QMessageBox.information(self, "Không có file .docx", "Không tìm thấy file .docx trong thư mục đã chọn.")
            return
        if csv_text == "ERROR":
            QMessageBox.critical(self, "Lỗi", "Đã xảy ra lỗi khi phân tích.")
            return
        self.csv_text = csv_text
        self.populate_table(result.get("summary_rows", []))
        self.log.append("Hoàn tất phân tích.")
        self.btn_csv.setEnabled(True)
        self.btn_xlsx.setEnabled(True)

    def populate_table(self, rows):
        self.table.setRowCount(0)
        for r, row in enumerate(rows):
            self.table.insertRow(r)
            for c, val in enumerate(row):
                item = QTableWidgetItem(val)
                self.table.setItem(r, c, item)
        self.table.resizeColumnsToContents()

    def on_export_csv(self):
        if not self.csv_text:
            QMessageBox.information(self, "Không có dữ liệu", "Chưa có dữ liệu để xuất.")
            return
        fname, _ = QFileDialog.getSaveFileName(self, "Lưu CSV", filter="CSV Files (*.csv)")
        if fname:
            with open(fname, "w", encoding="utf-8", newline="") as fh:
                fh.write(self.csv_text)
            QMessageBox.information(self, "Hoàn tất", f"Đã lưu CSV: {fname}")

    def on_export_xlsx(self):
        if not self.csv_text:
            QMessageBox.information(self, "Không có dữ liệu", "Chưa có dữ liệu để xuất.")
            return
        fname, _ = QFileDialog.getSaveFileName(self, "Lưu Excel", filter="Excel Files (*.xlsx)")
        if fname:
            wb = Workbook()
            ws = wb.active
            lines = self.csv_text.splitlines()
            for r, line in enumerate(lines, start=1):
                cols = [c for c in line.split(",")]
                for c, val in enumerate(cols, start=1):
                    ws.cell(row=r, column=c, value=val)
            wb.save(fname)
            QMessageBox.information(self, "Hoàn tất", f"Đã lưu Excel: {fname}")

def main():
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()
